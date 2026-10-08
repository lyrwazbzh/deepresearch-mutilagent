from __future__ import annotations

import html
import io
import json
import secrets
import tempfile
import threading
import time
import zipfile
from contextlib import nullcontext
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from typing import Any, Callable, Literal

from docx import Document
from docx.shared import Inches
from pydantic import BaseModel, ConfigDict, Field
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError

from src.agent.backends.audit import AuditService, current_trace_id
from src.agent.tools.web import SSRFDeniedError, SSRFGuard
from src.agent.identity import canonical_json, sha256_hex
from src.agent.schema.report import (
    ClaimSegment,
    ImageBlock,
    ListBlock,
    ParagraphBlock,
    ReportAst,
    TableBlock,
)
from src.agent.schema.state import Evidence, ResearchState
from src.agent.backends.persistence.database import SessionFactory
from src.agent.backends.persistence.models import (
    ExportJobRow,
    OutboxMessageRow,
    ResearchEventRow,
    ResearchTaskRow,
    ShareRow,
    utcnow,
)
from src.agent.backends.observability import Telemetry, TraceScope
from src.agent.reporting.reporting import CitationValidator, ReportDependencyError
from src.agent.errors import (
    ExportExpiredError,
    IdempotencyConflictError,
    InvalidArgumentError,
    InvalidStateError,
    NotFoundError,
    RateLimitedError,
)
from src.agent.task_service import new_id
from src.agent.backends.workspace import WorkspaceService


def _aware(value: datetime) -> datetime:
    return value if value.tzinfo is not None else value.replace(tzinfo=timezone.utc)


def _clean_text(value: str) -> str:
    """Keep the public JSON inert even if a future UI accidentally uses innerHTML."""
    cleaned = "".join(ch for ch in value if ch in "\n\t" or ord(ch) >= 32)
    return html.escape(cleaned, quote=True)


def _clean_json(value: Any) -> Any:
    if isinstance(value, str):
        return _clean_text(value)
    if isinstance(value, list):
        return [_clean_json(item) for item in value]
    if isinstance(value, dict):
        return {str(key): _clean_json(item) for key, item in value.items()}
    return value


class CreateShareRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    task_id: str = Field(min_length=1, max_length=64)
    report_version: int | None = Field(default=None, ge=1)
    title: str | None = Field(default=None, min_length=1, max_length=512)
    expires_at: datetime | None = None


class CreateExportRequest(BaseModel):
    model_config = ConfigDict(extra="forbid")
    format: Literal["docx", "pdf"]
    report_version: int | None = Field(default=None, ge=1)


class PublicRateLimiter:
    """Process-local bounded limiter, matching the existing in-process SSE admission model."""

    def __init__(self, *, requests: int = 60, window_seconds: int = 60) -> None:
        self.requests = max(1, requests)
        self.window_seconds = max(1, window_seconds)
        self._lock = threading.Lock()
        self._buckets: dict[str, tuple[int, int]] = {}

    def acquire(self, key: str, *, now: float | None = None) -> None:
        current = time.monotonic() if now is None else now
        bucket = int(current // self.window_seconds)
        with self._lock:
            stored_bucket, count = self._buckets.get(key, (bucket, 0))
            if stored_bucket != bucket:
                stored_bucket, count = bucket, 0
            if count >= self.requests:
                raise RateLimitedError("Public Share request rate exceeded.")
            self._buckets[key] = (stored_bucket, count + 1)
            if len(self._buckets) > 10_000:
                self._buckets = {
                    item_key: item
                    for item_key, item in self._buckets.items()
                    if item[0] >= bucket - 1
                }


@dataclass(frozen=True)
class LoadedReport:
    task: ResearchTaskRow
    state: ResearchState
    report: ReportAst
    evidence: list[Evidence]
    assets: dict[str, tuple[bytes, str, str]]


class FinalReportLoader:
    def __init__(
        self,
        session_factory: SessionFactory,
        workspace: WorkspaceService,
        citation_validator: CitationValidator,
    ) -> None:
        self.session_factory = session_factory
        self.workspace = workspace
        self.citation_validator = citation_validator

    def load(
        self,
        *,
        tenant_id: str,
        task_id: str,
        report_version: int,
        user_id: str | None = None,
        require_current_final: bool = False,
    ) -> LoadedReport:
        with self.session_factory() as session:
            statement = select(ResearchTaskRow).where(
                ResearchTaskRow.tenant_id == tenant_id,
                ResearchTaskRow.task_id == task_id,
                ResearchTaskRow.deleted_at.is_(None),
            )
            if user_id is not None:
                statement = statement.where(ResearchTaskRow.user_id == user_id)
            task = session.scalar(statement)
            if task is None:
                raise NotFoundError("Task was not found.")
            session.expunge(task)
        state = ResearchState.model_validate(task.latest_state)
        final = state.report.final
        if final is None:
            raise NotFoundError("A final report is not available for this task.")
        if require_current_final and final.version != report_version:
            raise InvalidArgumentError("The requested report version is not the current final report.")
        path = (
            final.workspace_path
            if final.version == report_version
            else f"reports/versions/{report_version}/report.json"
        )
        raw = self.workspace.read(tenant_id, task_id, path)
        if final.version == report_version and sha256_hex(raw) != final.content_hash:
            raise InvalidStateError("The final report content hash does not match ResearchState.")
        try:
            report = ReportAst.model_validate_json(raw)
        except ValueError as exc:
            raise InvalidStateError("The final report cannot be parsed.") from exc
        if report.report_version != report_version:
            raise InvalidStateError("The stored report version does not match the requested version.")
        citations = {item.evidence_id for item in report.citations}
        evidence = [item for item in state.evidence_pool if item.evidence_id in citations]
        if len({item.evidence_id for item in evidence}) != len(citations):
            raise InvalidStateError("The final report references unavailable Evidence.")
        validation = self.citation_validator.validate(
            report,
            state,
            tenant_id=tenant_id,
            user_id=task.user_id,
            task_id=task_id,
        )
        if not validation.valid:
            raise InvalidStateError("The final report failed deterministic citation validation.")
        assets: dict[str, tuple[bytes, str, str]] = {}
        for section in report.sections:
            for block in section.blocks:
                if not isinstance(block, ImageBlock) or block.workspace_path in assets:
                    continue
                metadata = self.workspace.metadata(tenant_id, task_id, block.workspace_path)
                payload = self.workspace.read(tenant_id, task_id, block.workspace_path)
                if metadata.content_hash != block.content_hash:
                    raise InvalidStateError("A report image content hash does not match Workspace.")
                assets[block.workspace_path] = (
                    payload,
                    metadata.mime_type,
                    metadata.content_hash,
                )
        return LoadedReport(task, state, report, evidence, assets)


class ShareManager:
    def __init__(
        self,
        session_factory: SessionFactory,
        workspace: WorkspaceService,
        loader: FinalReportLoader,
        audit: AuditService,
        *,
        default_days: int = 30,
        max_days: int = 90,
        public_base_url: str = "",
        url_validator: Callable[[str], str] | None = None,
    ) -> None:
        self.session_factory = session_factory
        self.workspace = workspace
        self.loader = loader
        self.audit = audit
        self.default_days = default_days
        self.max_days = max_days
        self.public_base_url = public_base_url.rstrip("/")
        self.url_validator = url_validator or SSRFGuard().validate

    def create(
        self, tenant_id: str, user_id: str, request: CreateShareRequest
    ) -> dict[str, object]:
        with self.session_factory() as session:
            task = session.scalar(
                select(ResearchTaskRow).where(
                    ResearchTaskRow.tenant_id == tenant_id,
                    ResearchTaskRow.user_id == user_id,
                    ResearchTaskRow.task_id == request.task_id,
                    ResearchTaskRow.deleted_at.is_(None),
                )
            )
            if task is None:
                raise NotFoundError("Task was not found.")
            state = ResearchState.model_validate(task.latest_state)
            if state.report.final is None:
                raise NotFoundError("A final report is not available for this task.")
            version = request.report_version or state.report.final.version
        loaded = self.loader.load(
            tenant_id=tenant_id,
            user_id=user_id,
            task_id=request.task_id,
            report_version=version,
            require_current_final=True,
        )
        now = utcnow()
        expires_at = request.expires_at or now + timedelta(days=self.default_days)
        expires_at = _aware(expires_at)
        if expires_at <= now or expires_at > now + timedelta(days=self.max_days):
            raise InvalidArgumentError(
                f"Share expiry must be in the future and no more than {self.max_days} days."
            )
        share_id = new_id("share")
        token = secrets.token_urlsafe(32)
        token_hash = sha256_hex(token)
        manifest: list[dict[str, object]] = []
        asset_ids: dict[str, str] = {}
        for path, (payload, mime_type, content_hash) in loaded.assets.items():
            asset_id = f"asset_{sha256_hex(f'{path}|{content_hash}')[:24]}"
            suffix = _safe_asset_suffix(mime_type)
            copied_path = f"reports/shares/{share_id}/assets/{asset_id}{suffix}"
            self.workspace.publish_bytes(
                tenant_id=tenant_id,
                task_id=request.task_id,
                path=copied_path,
                data=payload,
                mime_type=mime_type,
                operation_key=sha256_hex(f"share:{share_id}:asset:{asset_id}"),
            )
            asset_ids[path] = asset_id
            manifest.append(
                {
                    "asset_id": asset_id,
                    "media_type": mime_type,
                    "content_hash": content_hash,
                }
            )
        snapshot = self._snapshot(loaded, asset_ids, manifest, now)
        snapshot_path = f"reports/shares/{share_id}/snapshot.json"
        self.workspace.publish_bytes(
            tenant_id=tenant_id,
            task_id=request.task_id,
            path=snapshot_path,
            data=canonical_json(snapshot),
            mime_type="application/json",
            operation_key=sha256_hex(f"share:{share_id}:snapshot"),
        )
        title = _clean_text(request.title or loaded.report.title)
        try:
            with self.session_factory() as session, session.begin():
                task = session.scalar(
                    select(ResearchTaskRow)
                    .where(
                        ResearchTaskRow.tenant_id == tenant_id,
                        ResearchTaskRow.user_id == user_id,
                        ResearchTaskRow.task_id == request.task_id,
                        ResearchTaskRow.deleted_at.is_(None),
                    )
                    .with_for_update()
                )
                if task is None:
                    raise NotFoundError("Task was not found.")
                current = ResearchState.model_validate(task.latest_state).report.final
                if current is None or current.version != version:
                    raise InvalidStateError("The final report changed while the Share was created.")
                session.add(
                    ShareRow(
                        share_id=share_id,
                        tenant_id=tenant_id,
                        task_id=request.task_id,
                        report_version=version,
                        snapshot_workspace_path=snapshot_path,
                        public_token_hash=token_hash,
                        title=title,
                        expires_at=expires_at,
                        created_at=now,
                    )
                )
        except IntegrityError as exc:
            raise InvalidStateError("The Share control record could not be committed.") from exc
        self._audit(tenant_id, user_id, "share.create", f"share:{share_id}")
        public_path = f"/api/v1/public/shares/{token}"
        return {
            "share_id": share_id,
            "public_url": f"{self.public_base_url}{public_path}",
            "expires_at": expires_at.isoformat(),
        }

    def list_owned(self, tenant_id: str, user_id: str) -> list[dict[str, object]]:
        now = utcnow()
        with self.session_factory() as session:
            rows = session.scalars(
                select(ShareRow)
                .join(ResearchTaskRow, ResearchTaskRow.task_id == ShareRow.task_id)
                .where(
                    ShareRow.tenant_id == tenant_id,
                    ResearchTaskRow.user_id == user_id,
                    ResearchTaskRow.deleted_at.is_(None),
                )
                .order_by(ShareRow.created_at.desc(), ShareRow.share_id)
            ).all()
            return [self._projection(row, now) for row in rows]

    def revoke(self, tenant_id: str, user_id: str, share_id: str) -> None:
        with self.session_factory() as session, session.begin():
            row = session.scalar(
                select(ShareRow)
                .join(ResearchTaskRow, ResearchTaskRow.task_id == ShareRow.task_id)
                .where(
                    ShareRow.share_id == share_id,
                    ShareRow.tenant_id == tenant_id,
                    ResearchTaskRow.user_id == user_id,
                    ResearchTaskRow.deleted_at.is_(None),
                )
                .with_for_update()
            )
            if row is None:
                raise NotFoundError("Share was not found.")
            if row.revoked_at is None:
                row.revoked_at = utcnow()
        self._audit(tenant_id, user_id, "share.revoke", f"share:{share_id}")

    def read_public(self, token: str) -> tuple[ShareRow, dict[str, Any]]:
        row = self._active_by_token(token)
        raw = self.workspace.read(row.tenant_id, row.task_id, row.snapshot_workspace_path)
        snapshot = json.loads(raw)
        self.audit.append(
            tenant_id=row.tenant_id,
            actor_ref="public:share",
            action="share.read",
            resource_ref=f"share:{row.share_id}",
            result="succeeded",
            trace_id=current_trace_id(),
        )
        return row, snapshot

    def read_asset(self, token: str, asset_id: str) -> tuple[ShareRow, bytes, str, str]:
        row = self._active_by_token(token)
        raw = self.workspace.read(row.tenant_id, row.task_id, row.snapshot_workspace_path)
        snapshot = json.loads(raw)
        asset = next(
            (item for item in snapshot.get("assets", []) if item.get("asset_id") == asset_id),
            None,
        )
        if asset is None:
            raise NotFoundError("Share asset was not found.")
        path = (
            f"reports/shares/{row.share_id}/assets/{asset_id}"
            f"{_safe_asset_suffix(str(asset['media_type']))}"
        )
        payload = self.workspace.read(row.tenant_id, row.task_id, path)
        if sha256_hex(payload) != asset["content_hash"]:
            raise InvalidStateError("Share asset integrity check failed.")
        self.audit.append(
            tenant_id=row.tenant_id,
            actor_ref="public:share",
            action="share.asset.read",
            resource_ref=f"share:{row.share_id}:asset:{asset_id}",
            result="succeeded",
            trace_id=current_trace_id(),
        )
        return row, payload, str(asset["media_type"]), str(asset["content_hash"])

    def _active_by_token(self, token: str) -> ShareRow:
        if len(token) < 32 or len(token) > 256:
            raise NotFoundError("Share was not found.")
        digest = sha256_hex(token)
        with self.session_factory() as session:
            row = session.scalar(select(ShareRow).where(ShareRow.public_token_hash == digest))
            if (
                row is None
                or row.revoked_at is not None
                or _aware(row.expires_at) <= utcnow()
            ):
                raise NotFoundError("Share was not found.")
            session.expunge(row)
            return row

    def _snapshot(
        self,
        loaded: LoadedReport,
        asset_ids: dict[str, str],
        manifest: list[dict[str, object]],
        created_at: datetime,
    ) -> dict[str, Any]:
        public_report = loaded.report.model_dump(mode="json")
        for section in public_report["sections"]:
            for block in section["blocks"]:
                if block["type"] == "image":
                    block["asset_id"] = asset_ids[str(block.pop("workspace_path"))]
                    block.pop("content_hash", None)
        public_report["citations"] = [
            {
                "citation_id": citation.citation_id,
                **({"locator": citation.locator} if citation.locator else {}),
            }
            for citation in loaded.report.citations
        ]
        evidence = {item.evidence_id: item for item in loaded.evidence}
        source_cards: list[dict[str, object]] = []
        for citation in loaded.report.citations:
            item = evidence[citation.evidence_id]
            card: dict[str, object] = {
                "citation_id": citation.citation_id,
                "display_title": _clean_text(item.summary),
            }
            locator = citation.locator or _source_locator(item)
            if locator:
                card["locator"] = _clean_text(locator)
            if item.source_type == "web" and item.source_uri:
                try:
                    card["public_url"] = self.url_validator(item.source_uri)
                except (ValueError, SSRFDeniedError):
                    pass
            source_cards.append(card)
        return {
            "schema_version": 1,
            "title": _clean_text(loaded.report.title),
            "report_version": loaded.report.report_version,
            "public_report": _clean_json(public_report),
            "source_cards": source_cards,
            "assets": manifest,
            "created_at": created_at.isoformat(),
        }

    @staticmethod
    def _projection(row: ShareRow, now: datetime) -> dict[str, object]:
        status = (
            "revoked"
            if row.revoked_at is not None
            else "expired"
            if _aware(row.expires_at) <= now
            else "active"
        )
        return {
            "share_id": row.share_id,
            "task_id": row.task_id,
            "report_version": row.report_version,
            "title": row.title,
            "status": status,
            "expires_at": _aware(row.expires_at).isoformat(),
            "created_at": _aware(row.created_at).isoformat(),
        }

    def _audit(self, tenant_id: str, user_id: str, action: str, resource: str) -> None:
        self.audit.append(
            tenant_id=tenant_id,
            actor_ref=self.audit.user_actor(user_id),
            action=action,
            resource_ref=resource,
            result="succeeded",
            trace_id=current_trace_id(),
        )


class ReportRenderer:
    """Deterministic, offline renderer for the fixed ReportAst and resolved assets."""

    def render(
        self,
        report: ReportAst,
        evidence: list[Evidence],
        assets: dict[str, tuple[bytes, str, str]],
        format: Literal["docx", "pdf"],
    ) -> bytes:
        with tempfile.TemporaryDirectory(prefix="deepresearch-export-"):
            if format == "docx":
                return self._docx(report, evidence, assets)
            return self._pdf(report, evidence, assets)

    @staticmethod
    def _segments(segments: list[Any]) -> str:
        parts: list[str] = []
        for segment in segments:
            parts.append(segment.text)
            if isinstance(segment, ClaimSegment):
                parts.extend(f" [{citation}]" for citation in segment.citation_ids)
        return "".join(parts)

    def _docx(
        self,
        report: ReportAst,
        evidence: list[Evidence],
        assets: dict[str, tuple[bytes, str, str]],
    ) -> bytes:
        document = Document()
        props = document.core_properties
        props.title = report.title
        props.author = "DeepResearch"
        props.created = datetime(2000, 1, 1)
        props.modified = datetime(2000, 1, 1)
        document.add_heading(report.title, level=0)
        for section in report.sections:
            document.add_heading(section.heading, level=min(section.level, 9))
            for block in section.blocks:
                if isinstance(block, ParagraphBlock):
                    document.add_paragraph(self._segments(block.segments))
                elif isinstance(block, ListBlock):
                    style = "List Number" if block.ordered else "List Bullet"
                    for item in block.items:
                        document.add_paragraph(self._segments(item.segments), style=style)
                elif isinstance(block, TableBlock):
                    if block.caption:
                        document.add_paragraph(block.caption)
                    table = document.add_table(rows=1, cols=len(block.columns))
                    for index, value in enumerate(block.columns):
                        table.rows[0].cells[index].text = value
                    for row in block.rows:
                        cells = table.add_row().cells
                        for index, value in enumerate(row):
                            cells[index].text = value
                    if block.citation_ids:
                        document.add_paragraph(" ".join(f"[{x}]" for x in block.citation_ids))
                elif isinstance(block, ImageBlock):
                    payload = assets[block.workspace_path][0]
                    try:
                        document.add_picture(io.BytesIO(payload), width=Inches(6.0))
                    except Exception:
                        document.add_paragraph(f"[Image: {block.alt_text}]")
                    if block.caption:
                        document.add_paragraph(block.caption)
                    if block.citation_ids:
                        document.add_paragraph(" ".join(f"[{x}]" for x in block.citation_ids))
        document.add_page_break()
        document.add_heading("Sources", level=1)
        evidence_by_id = {item.evidence_id: item for item in evidence}
        for citation in report.citations:
            source = evidence_by_id[citation.evidence_id]
            suffix = f" — {citation.locator}" if citation.locator else ""
            document.add_paragraph(
                f"[{citation.citation_id}] {source.summary}{suffix}"
            )
        stream = io.BytesIO()
        document.save(stream)
        return _normalize_zip(stream.getvalue())

    def _pdf(
        self,
        report: ReportAst,
        evidence: list[Evidence],
        assets: dict[str, tuple[bytes, str, str]],
    ) -> bytes:
        try:
            from reportlab.lib.pagesizes import A4  # type: ignore[import-untyped]
            from reportlab.pdfbase import pdfmetrics  # type: ignore[import-untyped]
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont  # type: ignore[import-untyped]
            from reportlab.pdfgen import canvas  # type: ignore[import-untyped]
        except ImportError as exc:  # pragma: no cover - packaging contract
            raise RuntimeError("The export image is missing reportlab.") from exc
        del assets
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
        stream = io.BytesIO()
        page_width, page_height = A4
        output = canvas.Canvas(
            stream,
            pagesize=A4,
            invariant=1,
            pageCompression=1,
            initialFontName="STSong-Light",
        )
        output.setTitle(report.title)
        y = page_height - 48

        def line(text: str, size: int = 10, indent: int = 0) -> None:
            nonlocal y
            chunks = _wrap_text(text, max(20, 78 - indent * 4)) or [""]
            for chunk in chunks:
                if y < 48:
                    output.showPage()
                    y = page_height - 48
                output.setFont("STSong-Light", size)
                output.drawString(48 + indent * 14, y, chunk)
                y -= size + 5

        line(report.title, 18)
        y -= 8
        for section in report.sections:
            line(section.heading, max(11, 17 - section.level))
            for block in section.blocks:
                if isinstance(block, ParagraphBlock):
                    line(self._segments(block.segments))
                elif isinstance(block, ListBlock):
                    for index, item in enumerate(block.items, 1):
                        marker = f"{index}." if block.ordered else "•"
                        line(f"{marker} {self._segments(item.segments)}", indent=1)
                elif isinstance(block, TableBlock):
                    if block.caption:
                        line(block.caption, 9)
                    line(" | ".join(block.columns), 9)
                    for row in block.rows:
                        line(" | ".join(row), 9)
                    if block.citation_ids:
                        line(" ".join(f"[{x}]" for x in block.citation_ids), 8)
                elif isinstance(block, ImageBlock):
                    line(f"[Image: {block.alt_text}]", 9)
                    if block.caption:
                        line(block.caption, 9)
                    if block.citation_ids:
                        line(" ".join(f"[{x}]" for x in block.citation_ids), 8)
            y -= 4
        line("Sources", 14)
        evidence_by_id = {item.evidence_id: item for item in evidence}
        for citation in report.citations:
            evidence_item = evidence_by_id[citation.evidence_id]
            suffix = f" — {citation.locator}" if citation.locator else ""
            line(f"[{citation.citation_id}] {evidence_item.summary}{suffix}", 9)
        output.save()
        return stream.getvalue()


class ExportManager:
    def __init__(
        self,
        session_factory: SessionFactory,
        workspace: WorkspaceService,
        loader: FinalReportLoader,
        audit: AuditService,
    ) -> None:
        self.session_factory = session_factory
        self.workspace = workspace
        self.loader = loader
        self.audit = audit

    @staticmethod
    def request_hash(request: CreateExportRequest) -> str:
        semantics = {
            "format": request.format,
            "report_version": (
                {"kind": "explicit", "value": request.report_version}
                if request.report_version is not None
                else {"kind": "omitted"}
            ),
        }
        return sha256_hex(canonical_json(semantics))

    def create(
        self,
        tenant_id: str,
        user_id: str,
        task_id: str,
        idempotency_key: str,
        request: CreateExportRequest,
    ) -> tuple[dict[str, object], bool]:
        digest = self.request_hash(request)
        try:
            with self.session_factory() as session, session.begin():
                existing = session.scalar(
                    select(ExportJobRow).where(
                        ExportJobRow.tenant_id == tenant_id,
                        ExportJobRow.task_id == task_id,
                        ExportJobRow.idempotency_key == idempotency_key,
                    )
                )
                if existing is not None:
                    if existing.request_hash != digest:
                        raise IdempotencyConflictError(
                            "Idempotency-Key was replayed with different export semantics."
                        )
                    return self._projection(existing), True
                task = session.scalar(
                    select(ResearchTaskRow)
                    .where(
                        ResearchTaskRow.tenant_id == tenant_id,
                        ResearchTaskRow.user_id == user_id,
                        ResearchTaskRow.task_id == task_id,
                        ResearchTaskRow.deleted_at.is_(None),
                    )
                    .with_for_update()
                )
                if task is None:
                    raise NotFoundError("Task was not found.")
                final = ResearchState.model_validate(task.latest_state).report.final
                if final is None:
                    raise NotFoundError("A final report is not available for this task.")
                version = request.report_version or final.version
                if version != final.version:
                    raise InvalidArgumentError(
                        "The requested report version is not the current final report."
                    )
                export_id = new_id("export")
                now = utcnow()
                row = ExportJobRow(
                    export_id=export_id,
                    tenant_id=tenant_id,
                    task_id=task_id,
                    report_version=version,
                    format=request.format,
                    status="pending",
                    idempotency_key=idempotency_key,
                    request_hash=digest,
                    attempt_count=0,
                    created_at=now,
                    updated_at=now,
                )
                session.add(row)
                session.add(
                    OutboxMessageRow(
                        message_id=new_id("outbox"),
                        tenant_id=tenant_id,
                        kind="report.export",
                        aggregate_id=export_id,
                        payload_json={"export_id": export_id},
                        idempotency_key=f"export:{export_id}:attempt:0",
                        available_at=now,
                    )
                )
                session.flush()
                result = self._projection(row)
        except IntegrityError:
            with self.session_factory() as session:
                existing = session.scalar(
                    select(ExportJobRow).where(
                        ExportJobRow.tenant_id == tenant_id,
                        ExportJobRow.task_id == task_id,
                        ExportJobRow.idempotency_key == idempotency_key,
                    )
                )
                if existing is None or existing.request_hash != digest:
                    raise IdempotencyConflictError(
                        "Idempotency-Key was replayed with different export semantics."
                    )
                return self._projection(existing), True
        self.audit.append(
            tenant_id=tenant_id,
            actor_ref=self.audit.user_actor(user_id),
            action="export.create",
            resource_ref=f"task:{task_id}:export:{result['export_id']}",
            result="succeeded",
            trace_id=current_trace_id(),
        )
        return result, False

    def get(
        self, tenant_id: str, user_id: str, task_id: str, export_id: str
    ) -> dict[str, object]:
        return self._projection(self._owned(tenant_id, user_id, task_id, export_id))

    def content(
        self, tenant_id: str, user_id: str, task_id: str, export_id: str
    ) -> tuple[bytes, str, str, str]:
        row = self._owned(tenant_id, user_id, task_id, export_id)
        if row.status != "completed" or row.workspace_path is None:
            raise InvalidStateError("The export is not completed.")
        if row.expires_at is None or _aware(row.expires_at) <= utcnow():
            raise ExportExpiredError("The export has expired.")
        metadata = self.workspace.metadata(tenant_id, task_id, row.workspace_path)
        payload = self.workspace.read(tenant_id, task_id, row.workspace_path)
        self.audit.append(
            tenant_id=tenant_id,
            actor_ref=self.audit.user_actor(user_id),
            action="export.download",
            resource_ref=f"task:{task_id}:export:{export_id}",
            result="succeeded",
            trace_id=current_trace_id(),
        )
        return payload, metadata.mime_type, metadata.content_hash, row.format

    def _owned(
        self, tenant_id: str, user_id: str, task_id: str, export_id: str
    ) -> ExportJobRow:
        with self.session_factory() as session:
            row = session.scalar(
                select(ExportJobRow)
                .join(ResearchTaskRow, ResearchTaskRow.task_id == ExportJobRow.task_id)
                .where(
                    ExportJobRow.export_id == export_id,
                    ExportJobRow.tenant_id == tenant_id,
                    ExportJobRow.task_id == task_id,
                    ResearchTaskRow.user_id == user_id,
                    ResearchTaskRow.deleted_at.is_(None),
                )
            )
            if row is None:
                raise NotFoundError("Export was not found.")
            session.expunge(row)
            return row

    @staticmethod
    def _projection(row: ExportJobRow) -> dict[str, object]:
        result: dict[str, object] = {
            "export_id": row.export_id,
            "task_id": row.task_id,
            "report_version": row.report_version,
            "format": row.format,
            "status": row.status,
            "attempt_count": row.attempt_count,
            "created_at": _aware(row.created_at).isoformat(),
            "updated_at": _aware(row.updated_at).isoformat(),
        }
        for key in ("workspace_path", "error_code"):
            value = getattr(row, key)
            if value is not None:
                result[key] = value
        for key in ("completed_at", "expires_at"):
            value = getattr(row, key)
            if value is not None:
                result[key] = _aware(value).isoformat()
        return result


class ExportWorker:
    def __init__(
        self,
        session_factory: SessionFactory,
        workspace: WorkspaceService,
        loader: FinalReportLoader,
        renderer: ReportRenderer,
        *,
        max_attempts: int = 3,
        expires_days: int = 7,
        execution_timeout_seconds: int = 900,
        telemetry: Telemetry | None = None,
    ) -> None:
        self.session_factory = session_factory
        self.workspace = workspace
        self.loader = loader
        self.renderer = renderer
        self.max_attempts = max(1, max_attempts)
        self.expires_days = max(1, expires_days)
        self.execution_timeout_seconds = max(1, execution_timeout_seconds)
        self.telemetry = telemetry

    def run(self, export_id: str) -> bool:
        claim = self._claim(export_id)
        if claim is None:
            return False
        tenant_id, task_id, report_version, format, attempt = claim
        heartbeat_stop = threading.Event()
        heartbeat = threading.Thread(
            target=self._heartbeat_loop,
            args=(export_id, attempt, heartbeat_stop),
            name=f"export-heartbeat-{export_id}",
            daemon=True,
        )
        heartbeat.start()
        started = time.perf_counter()
        span_context = (
            self.telemetry.span(
                "report.export",
                scope=TraceScope(tenant_id=tenant_id, task_id=task_id),
                attributes={
                    "research.export_id": export_id,
                    "research.export_attempt": attempt,
                    "research.export_format": format,
                },
            )
            if self.telemetry is not None
            else nullcontext(None)
        )
        try:
            with span_context as span:
                loaded = self.loader.load(
                    tenant_id=tenant_id,
                    task_id=task_id,
                    report_version=report_version,
                )
                self.heartbeat(export_id, attempt)
                payload = self.renderer.render(
                    loaded.report, loaded.evidence, loaded.assets, format
                )
                self.heartbeat(export_id, attempt)
                mime_type = (
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                    if format == "docx"
                    else "application/pdf"
                )
                path = f"reports/exports/{export_id}.{format}"
                self.workspace.publish_bytes(
                    tenant_id=tenant_id,
                    task_id=task_id,
                    path=path,
                    data=payload,
                    mime_type=mime_type,
                    operation_key=sha256_hex(f"export:{export_id}:file"),
                )
                completed = self._complete(export_id, attempt, path)
                if span is not None:
                    span.set_attribute("research.export_status", "completed" if completed else "fenced")
                    span.set_attribute(
                        "research.renderer_duration_ms",
                        round((time.perf_counter() - started) * 1000, 3),
                    )
                return completed
        except (ReportDependencyError, ConnectionError, TimeoutError) as exc:
            self._fail(export_id, attempt, _error_code(exc), retryable=True)
            return False
        except Exception as exc:
            self._fail(export_id, attempt, _error_code(exc), retryable=False)
            return False
        finally:
            heartbeat_stop.set()
            heartbeat.join(timeout=2)

    def _heartbeat_loop(
        self, export_id: str, expected_attempt: int, stop: threading.Event
    ) -> None:
        interval = max(1.0, min(30.0, self.execution_timeout_seconds / 3))
        while not stop.wait(interval):
            if not self.heartbeat(export_id, expected_attempt):
                return

    def heartbeat(self, export_id: str, expected_attempt: int) -> bool:
        with self.session_factory() as session, session.begin():
            row = session.scalar(
                select(ExportJobRow).where(ExportJobRow.export_id == export_id).with_for_update()
            )
            if (
                row is None
                or row.status != "running"
                or row.attempt_count != expected_attempt
            ):
                return False
            row.updated_at = utcnow()
            return True

    def recover_once(self, *, now: datetime | None = None) -> int:
        now = now or utcnow()
        cutoff = now - timedelta(seconds=self.execution_timeout_seconds)
        recovered = 0
        with self.session_factory() as session, session.begin():
            rows = session.scalars(
                select(ExportJobRow)
                .where(
                    ExportJobRow.status == "running",
                    ExportJobRow.updated_at <= cutoff,
                )
                .with_for_update()
            ).all()
            for row in rows:
                attempt = row.attempt_count
                if attempt >= self.max_attempts:
                    row.status = "failed"
                    row.error_code = "export_timeout"
                    row.updated_at = now
                    _export_event(session, row, now)
                else:
                    row.status = "pending"
                    row.updated_at = now
                    session.add(
                        OutboxMessageRow(
                            message_id=new_id("outbox"),
                            tenant_id=row.tenant_id,
                            kind="report.export",
                            aggregate_id=row.export_id,
                            payload_json={"export_id": row.export_id},
                            idempotency_key=f"export:{row.export_id}:recovery:{attempt}",
                            available_at=now,
                        )
                    )
                recovered += 1
            pending = session.scalars(
                select(ExportJobRow)
                .where(
                    ExportJobRow.status == "pending",
                    ExportJobRow.updated_at <= cutoff,
                )
                .with_for_update()
            ).all()
            for row in pending:
                undispatched = session.scalar(
                    select(OutboxMessageRow.message_id).where(
                        OutboxMessageRow.kind == "report.export",
                        OutboxMessageRow.aggregate_id == row.export_id,
                        OutboxMessageRow.dispatched_at.is_(None),
                    )
                )
                if undispatched is not None:
                    continue
                bucket = int(now.timestamp()) // self.execution_timeout_seconds
                session.add(
                    OutboxMessageRow(
                        message_id=new_id("outbox"),
                        tenant_id=row.tenant_id,
                        kind="report.export",
                        aggregate_id=row.export_id,
                        payload_json={"export_id": row.export_id},
                        idempotency_key=f"export:{row.export_id}:redelivery:{bucket}",
                        available_at=now,
                    )
                )
                row.updated_at = now
                recovered += 1
        return recovered

    def _claim(
        self, export_id: str
    ) -> tuple[str, str, int, Literal["docx", "pdf"], int] | None:
        with self.session_factory() as session, session.begin():
            row = session.scalar(
                select(ExportJobRow).where(ExportJobRow.export_id == export_id).with_for_update()
            )
            if row is None or row.status != "pending":
                return None
            task = session.scalar(
                select(ResearchTaskRow).where(
                    ResearchTaskRow.task_id == row.task_id,
                    ResearchTaskRow.tenant_id == row.tenant_id,
                    ResearchTaskRow.deleted_at.is_(None),
                )
            )
            if task is None:
                row.status = "failed"
                row.error_code = "task_deleted"
                row.updated_at = utcnow()
                _export_event(session, row, row.updated_at)
                return None
            row.status = "running"
            row.attempt_count += 1
            row.error_code = None
            row.updated_at = utcnow()
            return (
                row.tenant_id,
                row.task_id,
                row.report_version,
                row.format,  # type: ignore[return-value]
                row.attempt_count,
            )

    def _complete(self, export_id: str, expected_attempt: int, path: str) -> bool:
        now = utcnow()
        with self.session_factory() as session, session.begin():
            row = session.scalar(
                select(ExportJobRow).where(ExportJobRow.export_id == export_id).with_for_update()
            )
            if (
                row is None
                or row.status != "running"
                or row.attempt_count != expected_attempt
            ):
                return False
            row.status = "completed"
            row.workspace_path = path
            row.completed_at = now
            row.expires_at = now + timedelta(days=self.expires_days)
            row.updated_at = now
            _export_event(session, row, now)
            return True

    def _fail(
        self, export_id: str, expected_attempt: int, error_code: str, *, retryable: bool
    ) -> None:
        now = utcnow()
        with self.session_factory() as session, session.begin():
            row = session.scalar(
                select(ExportJobRow).where(ExportJobRow.export_id == export_id).with_for_update()
            )
            if (
                row is None
                or row.status != "running"
                or row.attempt_count != expected_attempt
            ):
                return
            if retryable and row.attempt_count < self.max_attempts:
                row.status = "pending"
                row.error_code = error_code
                row.updated_at = now
                session.add(
                    OutboxMessageRow(
                        message_id=new_id("outbox"),
                        tenant_id=row.tenant_id,
                        kind="report.export",
                        aggregate_id=row.export_id,
                        payload_json={"export_id": row.export_id},
                        idempotency_key=f"export:{row.export_id}:retry:{row.attempt_count}",
                        available_at=now,
                    )
                )
            else:
                row.status = "failed"
                row.error_code = error_code
                row.updated_at = now
                _export_event(session, row, now)


def _export_event(session: object, row: ExportJobRow, now: datetime) -> None:
    event = ResearchEventRow(
        tenant_id=row.tenant_id,
        task_id=row.task_id,
        event_type="export_updated",
        payload_json={
            "event_id": 0,
            "task_id": row.task_id,
            "export_id": row.export_id,
            "status": row.status,
            "at": now.isoformat(),
        },
        created_at=now,
    )
    session.add(event)  # type: ignore[attr-defined]
    session.flush()  # type: ignore[attr-defined]
    event.payload_json = {**event.payload_json, "event_id": event.event_id}
    session.add(  # type: ignore[attr-defined]
        OutboxMessageRow(
            message_id=new_id("outbox"),
            tenant_id=row.tenant_id,
            kind="research.event.publish",
            aggregate_id=str(event.event_id),
            payload_json={"task_id": row.task_id, "event_id": event.event_id},
            idempotency_key=f"event:{event.event_id}",
            available_at=now,
        )
    )


def _normalize_zip(payload: bytes) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(payload), "r")
    target_stream = io.BytesIO()
    with source, zipfile.ZipFile(
        target_stream, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9
    ) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            target.writestr(info, source.read(name))
    return target_stream.getvalue()


def _wrap_text(value: str, width: int) -> list[str]:
    if not value:
        return [""]
    lines: list[str] = []
    for logical in value.splitlines() or [""]:
        while len(logical) > width:
            lines.append(logical[:width])
            logical = logical[width:]
        lines.append(logical)
    return lines


def _source_locator(item: Evidence) -> str | None:
    if item.source_type == "rag":
        return item.chunk_id
    if item.source_type == "sql":
        return "Structured data query"
    if item.source_type == "compute":
        return "Computed result"
    return None


def _safe_asset_suffix(mime_type: str) -> str:
    return {
        "image/png": ".png",
        "image/jpeg": ".jpg",
        "image/gif": ".gif",
        "image/webp": ".webp",
        "image/svg+xml": ".svg",
    }.get(mime_type.lower(), ".bin")


def _error_code(exc: Exception) -> str:
    name = type(exc).__name__
    normalized = "".join(
        ("_" + char.lower()) if char.isupper() else char for char in name
    ).lstrip("_")
    return normalized[:64] or "export_failed"
